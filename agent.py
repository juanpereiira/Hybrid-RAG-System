#IMPORTS
#----------------------------------------------------------------
print("Importing libraries...")
from huggingface_hub import InferenceClient
import os
import warnings
from io import BytesIO
from pathlib import Path

from pypdf import PdfReader
from docx import Document as DocxDocument

warnings.filterwarnings("ignore")
from getpass import getpass

import numpy as np
import pandas as pd
import faiss
import re

from collections import Counter

from sentence_transformers import SentenceTransformer
from transformers import pipeline
from rank_bm25 import BM25Okapi
from smolagents import CodeAgent, tool
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

print("Libraries imported successfully!")

documents = []
source_documents = []
chunks = []
chunk_texts = []
embeddings = None
vector_index = None
bm25 = None

#CHUNKING CONFIGS
#----------------------------------------------------------------
print("Chunking documents...")
text_splitter = RecursiveCharacterTextSplitter(
    chunk_size=300,
    chunk_overlap=50,
    separators=[
        "\n\n",
        "\n",
        ". ",
        " ",
        ""
    ]
)

#LOAD OPEN SOURCE EMBEDDING MODEL
#----------------------------------------------------------------
embedding_model_name = "sentence-transformers/all-MiniLM-L6-v2"

embedding_model = SentenceTransformer(
    embedding_model_name
)

#DOCUMENT UPLOAD
#----------------------------------------------------------------
def process_uploaded_document(file_bytes, filename):

    global documents
    global source_documents
    global chunks
    global chunk_texts
    global embeddings
    global vector_index
    global bm25

    extension = Path(filename).suffix.lower()

    # --------------------------------------------------------
    # EXTRACT TEXT
    # --------------------------------------------------------

    if extension == ".pdf":

        pdf = PdfReader(
            BytesIO(file_bytes)
        )

        pages = []

        for page in pdf.pages:

            text = page.extract_text()

            if text:
                pages.append(text)

        extracted_text = "\n\n".join(pages)

    elif extension == ".docx":

        docx_file = DocxDocument(
            BytesIO(file_bytes)
        )

        paragraphs = []

        for paragraph in docx_file.paragraphs:

            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        extracted_text = "\n\n".join(paragraphs)

    elif extension in [".txt", ".md"]:

        extracted_text = file_bytes.decode(
            "utf-8",
            errors="ignore"
        )

    else:

        raise ValueError(
            "Unsupported file type. "
            "Supported types: PDF, DOCX, TXT and MD."
        )

    extracted_text = extracted_text.strip()

    if not extracted_text:

        raise ValueError(
            "No readable text was found in the uploaded file."
        )

    # --------------------------------------------------------
    # CREATE DOCUMENT
    # --------------------------------------------------------

    document_id = (
        "UPLOAD_"
        + Path(filename).stem.upper()
        .replace(" ", "_")
    )

    documents = [
        {
            "id": document_id,
            "title": filename,
            "content": extracted_text
        }
    ]

    # --------------------------------------------------------
    # LANGCHAIN DOCUMENT
    # --------------------------------------------------------

    source_documents = [
        Document(
            page_content=extracted_text,
            metadata={
                "id": document_id,
                "title": filename
            }
        )
    ]

    # --------------------------------------------------------
    # CHUNKING
    # --------------------------------------------------------

    chunks = text_splitter.split_documents(source_documents)

    if not chunks:
        raise ValueError("The uploaded document could not be split into chunks.")

    # Give every chunk its own unique ID.
    for i, chunk in enumerate(chunks, start=1):
        chunk.metadata["chunk_id"] = f"DOC001_CHUNK_{i:04d}"

    print("Generated chunks:", len(chunks))
    # --------------------------------------------------------
    # EMBEDDINGS
    # --------------------------------------------------------

    chunk_texts = [
        chunk.page_content
        for chunk in chunks
    ]

    embeddings = embedding_model.encode(
        chunk_texts,
        normalize_embeddings=True,
        show_progress_bar=False
    )

    embeddings = np.asarray(
        embeddings,
        dtype="float32"
    )

    # --------------------------------------------------------
    # FAISS INDEX CREATION
    # --------------------------------------------------------

    dimension = embeddings.shape[1]

    vector_index = faiss.IndexFlatIP(
        dimension
    )

    vector_index.add(
        embeddings
    )

    # --------------------------------------------------------
    # BM25 INDEX CREATION
    # --------------------------------------------------------

    tokenized_documents = [
        re.findall(
            r"\b[a-zA-Z0-9]+\b",
            chunk.page_content.lower()
        )
        for chunk in chunks
    ]

    bm25 = BM25Okapi(
        tokenized_documents
    )

    # --------------------------------------------------------
    # RETURN INFORMATION
    # --------------------------------------------------------

    return {
        "filename": filename,
        "file_type": extension,
        "documents": len(documents),
        "chunks": len(chunks),
        "characters": len(extracted_text)
    }

# VECTOR RETRIEVAL FUNCTION
# ----------------------------------------------------------------

def vector_search(query, top_k=10):

    if vector_index is None:
        raise ValueError(
            "No document has been uploaded yet."
        )

    query_embedding = embedding_model.encode(
        [query],
        normalize_embeddings=True
    )

    query_embedding = np.asarray(
        query_embedding,
        dtype="float32" 
    )

    top_k = min(top_k, vector_index.ntotal)

    scores, indices = vector_index.search(
        query_embedding,
        top_k
    )

    results = []

    for score, index in zip(
        scores[0],
        indices[0]
    ):

        if index < 0:
            continue

        chunk = chunks[int(index)]

        results.append({
            "score": float(score),
            "id": chunk.metadata["chunk_id"],
            "title": chunk.metadata["title"],
            "text": chunk.page_content
        })

    return results

# KEYWORD RETRIEVAL FUNCTION
#func
def keyword_search(query, top_k=3):

    if bm25 is None:
        raise ValueError(
            "No document has been uploaded yet."
        )

    query_tokens = re.findall(
        r"\b[a-zA-Z0-9]+\b",
        query.lower()
    )

    scores = bm25.get_scores(
        query_tokens
    )

    top_indices = np.argsort(
        scores
    )[::-1][:top_k]

    results = []

    for index in top_indices:

        results.append({
            "score": float(scores[index]),
            "id": chunks[index].metadata["chunk_id"],
            "title": chunks[index].metadata["title"],
            "text": chunks[index].page_content
        })

    return results


#HYBRID SEARCH
#----------------------------------------------------------------

#Normalize Scores
def normalize_scores(scores):

    scores = np.array(
        scores,
        dtype=float
    )

    if np.max(scores) == np.min(scores):

        return np.ones(
            len(scores)
        )

    return (
        (scores - np.min(scores))
        /
        (np.max(scores) - np.min(scores))
    )

#Retrieval Function 

def hybrid_search(
    query,
    top_k=3,
    vector_weight=0.6,
    keyword_weight=0.4
):

    vector_results = vector_search(
        query,
        top_k=20
    )

    keyword_results = keyword_search(
        query,
        top_k=20
    )

    vector_scores = normalize_scores(
        [r["score"] for r in vector_results]
    )

    keyword_scores = normalize_scores(
        [r["score"] for r in keyword_results]
    )

    combined = {}

    # Add vector results

    for i, result in enumerate(
        vector_results
    ):

        combined[result["id"]] = {

            **result,

            "vector_score":
                vector_scores[i],

            "keyword_score":
                0
        }

    # Add keyword results

    for i, result in enumerate(
        keyword_results
    ):

        if result["id"] not in combined:

            combined[result["id"]] = {

                **result,

                "vector_score":
                    0,

                "keyword_score":
                    keyword_scores[i]
            }

        else:

            combined[
                result["id"]
            ]["keyword_score"] = keyword_scores[i]

    # Calculate final score

    for result in combined.values():

        result["hybrid_score"] = (

            vector_weight *
            result["vector_score"]

            +

            keyword_weight *
            result["keyword_score"]
        )

    results = sorted(
        combined.values(),
        key=lambda x:
            x["hybrid_score"],
        reverse=True
    )

    return results[:top_k]

#CREATING RETRIEVER TOOL

def retrieve_context(query):

    results = hybrid_search(
        query,
        top_k=4
    )

    if not results:
        return "No relevant information was found in the knowledge base."

    context = ""

    for i, result in enumerate(results, start=1):
        context += f"""
SOURCE {i}
Chunk ID: {result["id"]}
Title: {result["title"]}
Hybrid Score: {result["hybrid_score"]:.3f}

Content:
{result["text"]}


-------------------------
"""

    return context

#ADDING HUGGING FACE API KEY

from dotenv import load_dotenv
import os

load_dotenv()

HF_TOKEN = os.getenv("HF_TOKEN")

if not HF_TOKEN:
    raise ValueError("HF_TOKEN not found in .env file.")

print("HF_TOKEN loaded successfully.")

#CREATE RETRIEVAL TOOL

@tool #decorator tells the agent: Make this python function available as a tool that the agent can decide to call.
def knowledge_base_search(query: str) -> str:
    """
    Search the enterprise knowledge base.

    Args:
        query: A natural-language question or search query.

    Returns:
        Relevant passages from the knowledge base.
    """

    #triple quotes create a multiline string.
    #Agent uses the tool desc to understand what the tool does and when it should use it.
    return retrieve_context(query)

#Connect an OPEN WEIGHT LLM

from smolagents.agents import ChatMessage
from transformers import pipeline, AutoTokenizer, logging

logging.set_verbosity_error()
import os


class HFPipelineWrapper:

    def __init__(self, hf_pipeline):

        self.pipeline = hf_pipeline

        self.tokenizer = AutoTokenizer.from_pretrained(
            hf_pipeline.model.name_or_path,
            trust_remote_code=True
        )

    def generate(self, messages, **kwargs):

        formatted_messages = []

        # System instruction
        system_message = """
        You are a knowledge base assistant.

        Answer questions ONLY using the provided knowledge base.

        STRICT RULES:
        1. Use only the information provided in the knowledge base.
        2. Do not use web search.
        3. Do not use Wikipedia or external sources.
        4. Do not use your general knowledge.
        5. Do not invent information.
        6. If the knowledge base does not contain enough information,
        say:
        "The knowledge base does not contain enough information
        to answer this question."
        7. Give a concise, natural-language answer.
        """

        formatted_messages.append({
            "role": "system",
            "content": system_message
        })

        # Convert smolagents messages
        for msg in messages:

            if isinstance(msg, ChatMessage):

                content = msg.content

                # IMPORTANT:
                # smolagents may provide content as a list
                if isinstance(content, list):

                    text_parts = []

                    for item in content:

                        if isinstance(item, dict):

                            if "text" in item:
                                text_parts.append(
                                    str(item["text"])
                                )

                            elif "content" in item:
                                text_parts.append(
                                    str(item["content"])
                                )

                        else:
                            text_parts.append(
                                str(item)
                            )

                    content = "\n".join(text_parts)

                else:
                    content = str(content)

                formatted_messages.append({
                    "role": msg.role,
                    "content": content
                })

            elif isinstance(msg, dict):

                content = msg.get("content", "")

                # Handle list content
                if isinstance(content, list):

                    text_parts = []

                    for item in content:

                        if isinstance(item, dict):

                            if "text" in item:
                                text_parts.append(
                                    str(item["text"])
                                )

                            elif "content" in item:
                                text_parts.append(
                                    str(item["content"])
                                )

                        else:
                            text_parts.append(
                                str(item)
                            )

                    content = "\n".join(text_parts)

                else:
                    content = str(content)

                formatted_messages.append({
                    "role": msg.get("role", "user"),
                    "content": content
                })

        # Convert messages to Qwen chat format
        prompt = self.tokenizer.apply_chat_template(
            formatted_messages,
            tokenize=False,
            add_generation_prompt=True
        )

        # Generation parameters
        generation_params = {
            "max_new_tokens": kwargs.get(
                "max_tokens",
                128 #max number of new tokens qwen should generate
            ),
            "temperature": kwargs.get(
                "temperature",
                0.2 #controls randomness
            ),
            "do_sample": True,
            "return_full_text": False
        }

        # Generate response
        result = self.pipeline(
            prompt,
            **generation_params
        )

        generated_text = result[0]["generated_text"]

        # Make absolutely sure the output is a string
        if isinstance(generated_text, list):

            text_parts = []

            for item in generated_text:

                if isinstance(item, dict):
                    text_parts.append(
                        str(item.get("content", item.get("text", "")))
                    )
                else:
                    text_parts.append(
                        str(item)
                    )

            generated_text = "\n".join(text_parts)

        generated_text = str(generated_text).strip()

        # Remove Qwen end token
        generated_text = generated_text.replace(
            "<|im_end|>",
            ""
        ).strip()

        return ChatMessage(
            role="assistant",
            content=generated_text
        )


# Create Hugging Face pipeline

hf_pipeline = pipeline(
    "text-generation",
    model="Qwen/Qwen2.5-3B-Instruct",
    token=os.environ["HF_TOKEN"],
    trust_remote_code=True
)

# Wrap pipeline for smolagents

model = HFPipelineWrapper(
    hf_pipeline
)

print("Open-weight LLM connected successfully!")

# ============================================================
# RAG ASSISTANT
# ============================================================

def ask_assistant(question):

    # Retrieve information from the knowledge base
    context = retrieve_context(question)

    messages = [

        {
            "role": "system",
            "content": """
You are a knowledge base assistant.

Answer questions ONLY using the provided knowledge base.

STRICT RULES:
1. Use only the information provided in the knowledge base.
2. Do not use web search.
3. Do not use Wikipedia or external sources.
4. Do not use your general knowledge.
5. Do not invent information.
6. If the knowledge base does not contain enough information,
   say:
   "The knowledge base does not contain enough information
   to answer this question."
7. Give a concise, natural-language answer.
"""
        },

        {
            "role": "user",
            "content": f"""
Knowledge Base Context:

{context}

Question:

{question}

Answer using ONLY the knowledge base context.
"""
        }

    ]

    # Convert messages to Qwen format
    prompt = model.tokenizer.apply_chat_template(
        messages,
        tokenize=False,
        add_generation_prompt=True
    )

    # Generate answer
    result = hf_pipeline(
        prompt,
        max_new_tokens=128,
        temperature=0.2,
        do_sample=True,
        return_full_text=False
    )

    answer = result[0]["generated_text"]

    # Clean output
    answer = str(answer).strip()

    answer = answer.replace(
        "<|im_end|>",
        ""
    ).strip()

    return answer

def evaluate_session(results):

    prompt = f"""
Evaluate this RAG session using ONLY the provided context and answers.

Give scores from 1 to 5 for:

- Retrieval Relevance
- Answer Correctness
- Groundedness
- Completeness

Scoring guidelines:

Retrieval Relevance:
How relevant and useful is the retrieved context for answering the question?

Answer Correctness:
How accurately does the answer respond to the question using the provided context?

Groundedness:
How well is the answer supported by the retrieved context?

Completeness:
How completely does the answer address the question based on the available context?

Return ONLY these four lines:

Retrieval Relevance: X/5
Answer Correctness: X/5
Groundedness: X/5
Completeness: X/5

Session:
{results}
"""

    messages = [
        {
            "role": "system",
            "content": (
                "You are a strict RAG evaluator. "
                "Use only the provided session information. "
                "Do not use external knowledge."
            )
        },
        {
            "role": "user",
            "content": prompt
        }
    ]

    prompt = model.tokenizer.apply_chat_template(
        messages,
        tokenize=False,
        add_generation_prompt=True
    )

    result = hf_pipeline(
        prompt,
        max_new_tokens=64,
        temperature=0.1,
        do_sample=False,
        return_full_text=False
    )

    evaluation = str(
        result[0]["generated_text"]
    ).strip()

    # Extract the four LLM-generated scores
    scores = {}

    for line in evaluation.split("\n"):
        match = re.match(
            r"^(Retrieval Relevance|Answer Correctness|Groundedness|Completeness):\s*(\d+(?:\.\d+)?)/5",
            line.strip()
        )

        if match:
            metric = match.group(1)
            score = float(match.group(2))
            scores[metric] = score

    # Calculate Overall mathematically
    required_metrics = [
        "Retrieval Relevance",
        "Answer Correctness",
        "Groundedness",
        "Completeness"
    ]

    if all(metric in scores for metric in required_metrics):

        overall = sum(
            scores[metric]
            for metric in required_metrics
        ) / len(required_metrics)

        overall = round(overall, 2)

    else:
        overall = 0.0

    return (
        f"Retrieval Relevance: "
        f"{scores.get('Retrieval Relevance', 0):.2f}/5\n"

        f"Answer Correctness: "
        f"{scores.get('Answer Correctness', 0):.2f}/5\n"

        f"Groundedness: "
        f"{scores.get('Groundedness', 0):.2f}/5\n"

        f"Completeness: "
        f"{scores.get('Completeness', 0):.2f}/5\n"

        f"Overall: "
        f"{overall:.2f}/5"
    )

